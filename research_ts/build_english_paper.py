"""Build the separately authored English manuscript from archived evidence."""
import copy,csv,json,re,hashlib
from pathlib import Path
import numpy as np
from docx import Document
from docx.shared import Mm,Pt,RGBColor,Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from paper_fonts import normalize

ROOT=Path(__file__).resolve().parents[1];PAPER=ROOT/'paper';RESULTS=ROOT/'results/timeseries_v2'
FINAL=ROOT/'DHPF_Time_Series_Generation_English.docx'
MODELS=['gan','ddpm_v','icfm','otcfm','dhpf','drifting']
LABELS=dict(gan='GAN',ddpm_v='DDPM-v',icfm='I-CFM',otcfm='OT-CFM',dhpf='DHPF',drifting='Drifting-raw')
def read(p):return json.loads(p.read_text(encoding='utf-8'))
def fmt(r,k):return f"{r[k+'_mean']:.4f}\n±{r[k+'_std']:.4f}"

def prose(doc,text,style_name=None):
    """Typeset inline indices and powers without exposing source notation."""
    p=doc.add_paragraph(style=style_name)
    tokens=r'(p_data|S_B|y_B|x̂_B|σ_min|[TGa]θ|\^\([^)]*\))'
    for part in re.split(tokens,text):
        if not part:continue
        if part.startswith('^('):
            p.add_run(part[2:-1]).font.superscript=True
        elif '_' in part and re.fullmatch(tokens,part):
            base,sub=part.split('_',1)
            p.add_run(base).italic=True;p.add_run(sub).font.subscript=True
        elif re.fullmatch('[TGa]θ',part):
            p.add_run(part[0]).italic=True;p.add_run('θ').font.subscript=True
        else:p.add_run(part)
    return p

def style(doc,name,size,before=0,after=6,bold=False,line=1.12):
    s=doc.styles[name] if name in doc.styles else doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    s.font.name='Times New Roman';s.font.size=Pt(size);s.font.bold=bold;s.font.color.rgb=RGBColor(0,0,0)
    pf=s.paragraph_format;pf.space_before=Pt(before);pf.space_after=Pt(after);pf.line_spacing=line
    pf.first_line_indent=Pt(0);pf.widow_control=True
    for e in s.element.xpath('./w:pPr/w:pBdr'):e.getparent().remove(e)
    return s

def table(doc,caption,headers,rows,note):
    doc.add_paragraph(caption,'Table Caption')
    tab=doc.add_table(rows=1,cols=len(headers));tab.autofit=False
    total=9411;first=1900 if len(headers)>=5 else 2300
    widths=[first]+[(total-first)//(len(headers)-1)]*(len(headers)-1);widths[-1]+=total-sum(widths)
    pr=tab._tbl.tblPr
    for tag in ['tblW','tblInd','tblLayout','tblCellMar','tblBorders']:
        old=pr.find(qn('w:'+tag))
        if old is not None:pr.remove(old)
    for tag,attrs in [('tblW',dict(w=str(total),type='dxa')),('tblInd',dict(w='100',type='dxa')),('tblLayout',dict(type='fixed'))]:
        e=OxmlElement('w:'+tag)
        for k,v in attrs.items():e.set(qn('w:'+k),v)
        pr.append(e)
    margins=OxmlElement('w:tblCellMar')
    for side,val in [('top',70),('bottom',70),('start',100),('end',100)]:
        e=OxmlElement('w:'+side);e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa');margins.append(e)
    pr.append(margins)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','bottom','left','right','insideH','insideV']:
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'C5C5C5');borders.append(e)
    pr.append(borders)
    for col,w in zip(tab.columns,widths):col.width=Twips(w)
    for cells,values in [(tab.rows[0].cells,headers)]+[(tab.add_row().cells,r) for r in rows]:
        for cell,val,w in zip(cells,values,widths):
            cell.width=Twips(w);cell.vertical_alignment=1;cell.text=str(val)
            for p in cell.paragraphs:p.style='Table Text';p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for c in tab.rows[0].cells:
        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F2F4F6');c._tc.get_or_add_tcPr().append(sh)
        for p in c.paragraphs:
            for r in p.runs:r.bold=True
    repeat=OxmlElement('w:tblHeader');tab.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in tab.rows:row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    for row in tab.rows[:-1] if len(rows)<=3 else tab.rows[-2:]:
        for c in row.cells:
            for p in c.paragraphs:p.paragraph_format.keep_with_next=True
    if note:doc.add_paragraph(note,'Table Note')
    return tab

def main():
    s={(r['dataset'],r['model'],r['variant']):r for r in read(RESULTS/'summary.json')}
    refs=read(RESULTS/'real_reference.json');ridge={}
    for m in MODELS:
        name='ddpm' if m=='ddpm_v' else m
        vals=[read(RESULTS/'ridge_probe'/f'har__{name}__base__seed{k}.json')['test_macro_f1'] for k in range(5)]
        ridge[m]=dict(f1_mean=float(np.mean(vals)),f1_std=float(np.std(vals,ddof=1)))
    rr=read(RESULTS/'ridge_probe/real_matched.json')['test_macro_f1']
    doc=Document();sec=doc.sections[0];sec.page_width=Mm(210);sec.page_height=Mm(297)
    for side in ['top_margin','bottom_margin','left_margin','right_margin']:setattr(sec,side,Mm(22))
    sec.header_distance=Mm(10);sec.footer_distance=Mm(12)
    style(doc,'Normal',11,after=6).paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    style(doc,'Title',18,after=12,bold=True)
    for name,size,before in [('Heading 1',13,12),('Heading 2',11.5,9)]:
        style(doc,name,size,before=before,after=6,bold=True).paragraph_format.keep_with_next=True
    for name,size,after in [('Caption',9,8),('Table Caption',9.5,5),('Table Note',8.5,7),('Table Text',9,0),('Equation',10,7),('Reference',9,4),('Algorithm',9,9)]:
        style(doc,name,size,after=after,line=1.05 if name=='Table Text' else 1.10)
    doc.styles['Table Caption'].paragraph_format.keep_with_next=True
    doc.styles['Table Note'].paragraph_format.space_before=Pt(3)
    doc.styles['Reference'].paragraph_format.left_indent=Pt(15);doc.styles['Reference'].paragraph_format.first_line_indent=Pt(-15)
    header=sec.header.paragraphs[0];header.text='DHPF | Multivariate Time-Series Generation'
    for r in header.runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('666666')
    foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
    math=Document(ROOT/'DHPF_时序数据生成研究论文.docx')._element.xpath('.//m:oMath')
    assert len(math)==11
    def equation(n):
        p=doc.add_paragraph(style='Equation');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p._p.append(copy.deepcopy(math[n-1]));p.add_run(f'    ({n})');p.paragraph_format.keep_together=True
    generic='Mean ± sample standard deviation over five training seeds. '
    def add_table(key):
        if key in ['sine','lorenz']:
            n=1 if key=='sine' else 2
            keys=['sw2_squared','power_rmse','autocorrelation_rmse','cross_correlation_rmse','synthetic_test_forecast_mse']
            rows=[[LABELS[m]]+[fmt(s[key,m,'base'],k) for k in keys] for m in MODELS]
            real=[f"{refs[key][k]['mean']:.4f}\n±{refs[key][k]['std']:.4f}" for k in keys[:-1]]+[f"{s[key,'dhpf','base']['real_matched_test_forecast_mse_mean']:.4f}"]
            rows.append(['Real reference']+real)
            table(doc,f'Table {n}. {key.title()} generation. Lower values are preferable for every column.',
                  ['Method','SW₂²','Power\nRMSE','ACF\nRMSE','Corr₀\nRMSE','TSTR\nMSE'],rows,
                  generic+'The real reference uses five training-set subsamples for distribution metrics and one equally sized real sample for the probe; its variation is not generator-seed variation.')
        elif key=='har':
            keys=['sw2_squared','autocorrelation_rmse','cross_correlation_rmse','synthetic_test_macro_f1']
            rows=[[LABELS[m]]+[fmt(s['har',m,'base'],k) for k in keys]+[fmt(ridge[m],'f1')] for m in MODELS]
            rows.append(['Real reference']+[f"{refs['har'][k]['mean']:.4f}\n±{refs['har'][k]['std']:.4f}" for k in keys[:3]]+[f"{s['har','dhpf','base']['real_matched_test_macro_f1_mean']:.4f}",f'{rr:.4f}'])
            table(doc,'Table 3. Class-conditional HAR generation and downstream classification.',
                  ['Method','SW₂² ↓','ACF ↓','Corr₀ ↓','LR F1 ↑','Ridge F1 ↑'],rows,
                  generic+'Distribution metrics are averaged equally across six classes. F1 denotes macro-F1 on real test subjects. The real-reference distribution metrics include subject shift and sampling variability.')
        elif key=='ablations':
            variants=[('dhpf','base','Full DHPF'),('dhpf','fixed_full','Fixed centers'),('dhpf','radial','Radial only'),('dhpf','tangential','Solenoidal only'),('direct','base','MLP acceleration'),('augmented','base','Augmented ODE')]
            rows=[[label]+[fmt(s['sine',m,v],k) for k in ['sw2_squared','autocorrelation_rmse','cross_correlation_rmse','synthetic_test_forecast_mse']] for m,v,label in variants]
            table(doc,'Table 4. Sine component and dynamics controls under terminal matching.',
                  ['Model','SW₂² ↓','ACF ↓','Corr₀ ↓','TSTR MSE ↓'],rows,generic+'All controls use the same endpoint objective and update budget. Fixed centers remain trainable; their time-dependent offsets are disabled.')
        elif key=='timing':
            times={r['model']:r for r in read(RESULTS/'inference_timing.json') if r['dataset']=='sine'}
            rows=[[LABELS[m],f"{s['sine',m,'base']['update_seconds_mean']*1000:.2f}",str(dict(gan=1,ddpm_v=1000,icfm=64,otcfm=64,dhpf=64,drifting=1)[m]),f"{times[m]['median_seconds']*1000:.2f}"] for m in MODELS]
            table(doc,'Table 5. Computation on Sine with one RTX 5070 Ti.',
                  ['Method','Training\nms/update','Sampling\nqueries/passes','Generation\nms / 512 windows'],rows,
                  'Training time is the mean synchronized update time over five runs. Generation time is the median of three technical repeats from the seed-0 checkpoint, including prior generation and CPU output transfer. DHPF counts analytic field queries after shared source preparation; query costs differ across methods.')
        elif key=='datasets':
            table(doc,'Table B1. Dataset sizes and conditioning.',
                  ['Dataset','Window shape','Train / validation / test','Condition'],
                  [['Sine','128 × 3','10,000 / 2,000 / 2,000','None'],['Lorenz','128 × 3','10,000 / 2,000 / 2,000','None'],['UCI HAR','128 × 9','5,883 / 1,469 / 2,947','Six activities']],
                  'Normalization is fitted on the training split. HAR subjects are disjoint across splits.')
        elif key=='extra_metrics':
            rows=[[LABELS[m]]+[fmt(s[d,m,'base'],'mmd2_unbiased') for d in ['sine','lorenz','har']]+[fmt(s['har',m,'base'],'power_rmse')] for m in MODELS]
            table(doc,'Table C1. Additional distribution and power-spectrum metrics (lower is better).',
                  ['Method','Sine MMD²','Lorenz MMD²','HAR MMD²','HAR power\nRMSE'],rows,generic+'MMD² uses the U-statistic form, with its independence assumption discussed in Appendix B.')
        elif key=='har_classes':
            with (RESULTS/'per_class.csv').open(encoding='utf-8-sig') as f:raw=list(csv.DictReader(f))
            rows=[]
            for c,label in enumerate(['Walking','Upstairs','Downstairs','Sitting','Standing','Lying']):
                records=[r for r in raw if r['dataset']=='har' and r['model']=='dhpf' and r['variant']=='base' and int(r['condition'])==c]
                assert len(records)==5
                rows.append([label]+[f"{np.mean([float(r[k]) for r in records]):.4f}\n±{np.std([float(r[k]) for r in records],ddof=1):.4f}" for k in ['sw2_squared','autocorrelation_rmse','cross_correlation_rmse','power_rmse']])
            table(doc,'Table C2. DHPF results for individual HAR activities (lower is better).',
                  ['Activity','SW₂²','ACF RMSE','Corr₀ RMSE','Power RMSE'],rows,generic)
        else:raise ValueError(key)
    algorithms={
        'training':('Algorithm 1. Training DHPF',[
            'Input: dataset, source parameters θ, batch size B, K midpoint steps.',
            'Repeat for the prescribed number of updates:',
            '    Select c uniformly for a conditional task; otherwise omit c.',
            '    Sample B real windows y₁,…,y_B from the selected condition.',
            '    Draw B independent initial states using Equation (2).',
            '    Evaluate Gθ at all K grid times and K midpoints.',
            '    Integrate Equation (8) to obtain generated endpoints x̂₁,…,x̂_B.',
            '    Solve the assignment in Equation (7) with detached costs.',
            '    Evaluate the selected endpoint loss with its gradient graph.',
            '    Backpropagate through the solver; clip gradients; update θ with Adam.']),
        'sampling':('Algorithm 2. Sampling DHPF',[
            'Input: trained θ, sample count, optional condition c, K midpoint steps.',
            'Draw independent initial states using Equation (2).',
            'Evaluate Gθ at the grid times and midpoints for the chosen condition.',
            'Integrate Equation (8) without retaining a gradient graph.',
            'Reshape terminal positions into L-by-C windows and return them.'])}
    citations=[b for b in (PAPER/'manuscript_template.txt').read_text(encoding='utf-8').split('\n\n') if re.match(r'^\[\d+\]',b)]
    assert len(citations)==20
    text=(PAPER/'manuscript_english.txt').read_text(encoding='utf-8')
    for block in text.strip().split('\n\n'):
        if block.startswith('{{equation:'):equation(int(block[11:-2]))
        elif block.startswith('{{table:'):add_table(block[8:-2])
        elif block.startswith('{{figure:'):
            name=block[9:-2];p=doc.add_paragraph();p.paragraph_format.keep_with_next=True
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(PAPER/'figures'/f'{name}.png'),width=Mm(164))
            for e in p._p.xpath('.//wp:docPr'):e.set('descr',name.replace('_',' '))
        elif block.startswith('{{algorithm:'):
            key=block[12:-2];title,lines=algorithms[key];doc.add_paragraph(title,'Table Caption')
            p=prose(doc,'\n'.join(lines),'Algorithm');p.paragraph_format.keep_together=True
            pr=p._p.get_or_add_pPr();borders=OxmlElement('w:pBdr')
            for side in ['top','bottom']:
                e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:space'),'6');borders.append(e)
            pr.append(borders)
        elif block=='{{references}}':
            for b in citations:doc.add_paragraph(b.strip(),'Reference')
        elif block.startswith('### '):doc.add_paragraph(block[4:],'Heading 2')
        elif block.startswith('## '):
            p=doc.add_paragraph(block[3:],'Heading 1')
            if block.startswith('## References') or block.startswith('## Appendix A.'):
                p.paragraph_format.page_break_before=True
        elif block.startswith('# '):doc.add_paragraph(block[2:].replace(' for Multivariate','\nfor Multivariate'),'Title').alignment=WD_ALIGN_PARAGRAPH.CENTER
        elif re.match(r'^Figure \d+\.',block):prose(doc,block,'Caption')
        else:prose(doc,block)
    normalize(doc)
    # English manuscript styles are explicitly Roman, including headings.
    for st in doc.styles:
        if st.element.rPr is not None:
            for el in st.element.rPr.findall(qn('w:rFonts')):el.set(qn('w:hint'),'default')
    doc.core_properties.title='Dynamic Helmholtz Particle Flow for Multivariate Time-Series Generation'
    doc.core_properties.subject='English research manuscript';doc.core_properties.author='';doc.core_properties.comments=''
    doc.save(FINAL)
    assert '{{' not in '\n'.join(p.text for p in doc.paragraphs)
    assert len(doc._element.xpath('.//m:oMath'))==11 and len(doc.tables)==8 and len(doc.inline_shapes)==6
    metadata=dict(document=FINAL.name,sha256=hashlib.sha256(FINAL.read_bytes()).hexdigest(),
                  preset='narrative_proposal',header_pattern='memo_masthead',named_override='EnglishAcademic',
                  page_mm=[210,297],margins_mm=22,font='Times New Roman',body_pt=11,body_line_spacing=1.12,
                  equations=11,tables=8,figures=6,algorithms=2,references=20,
                  evidence='results/timeseries_v2/summary.json',source='paper/manuscript_english.txt')
    (PAPER/'english_build.json').write_text(json.dumps(metadata,indent=2),encoding='utf-8')
    print(json.dumps(metadata,indent=2))

if __name__=='__main__':main()
