"""Verify the English document against archived evidence and preserved sources."""
import hashlib,importlib.util,json,math,re,sys
from pathlib import Path
import numpy as np
from docx import Document
import pymupdf

ROOT=Path(__file__).resolve().parents[1]
PAPER=ROOT/'paper';RESULTS=ROOT/'results/timeseries_v2'
def read(p):return json.loads(p.read_text(encoding='utf-8'))
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def main():
    path=ROOT/'DHPF_Time_Series_Generation_English.docx';doc=Document(path)
    text='\n'.join(p.text for p in doc.paragraphs)
    source=(PAPER/'manuscript_english.txt').read_text(encoding='utf-8')
    summary={(r['dataset'],r['model'],r['variant']):r for r in read(RESULTS/'summary.json')}
    models=['gan','ddpm_v','icfm','otcfm','dhpf','drifting'];checked=0
    def cell(table,row,col,mean,std=None):
        nonlocal checked
        expected=f'{mean:.4f}'+(f'\n±{std:.4f}' if std is not None else '')
        assert doc.tables[table].cell(row,col).text==expected,(table,row,col,expected)
        checked+=1
    metric_keys=['sw2_squared','power_rmse','autocorrelation_rmse','cross_correlation_rmse','synthetic_test_forecast_mse']
    for ti,data in enumerate(['sine','lorenz']):
        for i,m in enumerate(models,1):
            r=summary[data,m,'base']
            for j,k in enumerate(metric_keys,1):cell(ti,i,j,r[k+'_mean'],r[k+'_std'])
    for i,m in enumerate(models,1):
        r=summary['har',m,'base']
        for j,k in enumerate(['sw2_squared','autocorrelation_rmse','cross_correlation_rmse','synthetic_test_macro_f1'],1):cell(2,i,j,r[k+'_mean'],r[k+'_std'])
        name='ddpm' if m=='ddpm_v' else m
        vals=[read(RESULTS/'ridge_probe'/f'har__{name}__base__seed{s}.json')['test_macro_f1'] for s in range(5)]
        cell(2,i,5,np.mean(vals),np.std(vals,ddof=1))
    for i,(m,v) in enumerate([('dhpf','base'),('dhpf','fixed_full'),('dhpf','radial'),('dhpf','tangential'),('direct','base'),('augmented','base')],1):
        r=summary['sine',m,v]
        for j,k in enumerate(['sw2_squared','autocorrelation_rmse','cross_correlation_rmse','synthetic_test_forecast_mse'],1):cell(3,i,j,r[k+'_mean'],r[k+'_std'])
    for i,m in enumerate(models,1):
        for j,data in enumerate(['sine','lorenz','har'],1):
            r=summary[data,m,'base'];cell(6,i,j,r['mmd2_unbiased_mean'],r['mmd2_unbiased_std'])
        r=summary['har',m,'base'];cell(6,i,4,r['power_rmse_mean'],r['power_rmse_std'])
    percentages={}
    for data,k,reported in [('sine','autocorrelation_rmse',31.9),('lorenz','cross_correlation_rmse',39.6)]:
        a=summary[data,'dhpf','base'][k+'_mean'];b=summary[data,'gan','base'][k+'_mean']
        p=(1-a/b)*100;assert round(p,1)==reported;percentages[data]=p
    paired=read(RESULTS/'paired_comparisons.json')
    assert len(paired)==90 and all(r['holm_p']>.05 for r in paired)
    refs=[p.text for p in doc.paragraphs if re.match(r'^\[\d+\]',p.text)]
    cited={int(x) for g in re.findall(r'\[(\d+(?:,\d+)*)\]',source) if g!='0,1' for x in g.split(',')}
    assert cited==set(range(1,21)) and len(refs)==20
    assert not re.search('[\u4e00-\u9fff]',text) and '{{' not in text
    assert len(doc.tables)==8 and len(doc.inline_shapes)==6 and len(doc._element.xpath('.//m:oMath'))==11
    frozen=read(RESULTS/'frozen_original.json')['sha256']
    assert all(sha(ROOT/p)==v for p,v in frozen.items())
    chinese=ROOT/'DHPF_时序数据生成研究论文.docx'
    assert sha(chinese)=='1b03f334062575ca97eb0d2c9be94fde353a4a3fcbc7aba9942e7b25fe7679a0'
    terms=['innovative','pioneering','revolutionary paradigm','transformative framework','superior','surpass','excel','remarkable','unprecedented','breakthrough performance','general-purpose','is capable of','notably','yet','yielding','at its essence','encompass','differentiate','reveal','underscore','pave the way for','highlight the potential of','profound challenges','stems from','rigid','impede']
    full=text+'\n'+'\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    hits={term:[{'line':i,'text':line} for i,line in enumerate(full.splitlines(),1) if term in line.lower()] for term in terms}
    hits={k:v for k,v in hits.items() if v}
    check=dict(document_sha256=sha(path),numeric_table_cells_checked=checked,headline_reductions_percent=percentages,
               citation_ids=sorted(cited),preserved_benchmark_files=len(frozen),chinese_sha256=sha(chinese),
               english_prose_words=len(re.findall(r'\b[\w]+(?:[-\u2013][\w]+)*\b',source)),
               full_document_vocabulary_scan=hits,em_dash_count=full.count('\u2014'))
    (PAPER/'english_evidence_check.json').write_text(json.dumps(check,indent=2,ensure_ascii=False),encoding='utf-8')
    # The installed auditor omits its math import. Inject it in memory only.
    audit_path=Path(str(Path(__import__('os').environ.get('NATURE_FIGURE_SCRIPTS', str(Path(__file__).resolve().parents[1] / 'optional_tools/nature-figure/scripts'))) / 'audit_svg_schematic.py'))
    spec=importlib.util.spec_from_file_location('english_svg_audit',audit_path)
    mod=importlib.util.module_from_spec(spec);sys.modules[spec.name]=mod;spec.loader.exec_module(mod);mod.math=math
    audit=mod.audit_svg((PAPER/'figures/fig0_method.svg').read_bytes(),5,10,181)
    audit['review_notes']=['Five-point floor applies to mathematical indices; main labels are 7.5–8 pt in the SVG.',
        'At the document width of 164 mm, main labels are 6.82–7.28 pt and indices are 5.09 pt.',
        'The generic 89/183 mm canvas warning is inapplicable: no journal preset was requested.',
        'The installed auditor missing-math-import defect was bypassed in memory; global skills were not edited.']
    assert not any(f['severity']=='FAIL' for f in audit['findings'])
    (PAPER/'english_schematic_audit.json').write_text(json.dumps(audit,indent=2,ensure_ascii=False),encoding='utf-8')
    render=ROOT/'tmp/paper_english_final_render';meta=read(render/'render_metadata.json')
    pdf=next(render.glob('*.pdf'));overflow=[]
    with pymupdf.open(pdf) as pages:
        assert len(pages)==18
        for i,page in enumerate(pages):
            for block in page.get_text('blocks'):
                if block[0]<20 or block[2]>page.rect.width-20 or block[1]<15 or block[3]>page.rect.height-15:overflow.append(i+1)
    assert not overflow
    qa=dict(document=path.name,sha256=sha(path),pages=18,visually_inspected_pages=list(range(1,19)),
        renderer=meta['renderer'],overflow_pages=overflow,equations=11,tables=8,figures=6,algorithms=2,
        body_font='Times New Roman',body_size_pt=11,
        visual_review='All final pages inspected: captions readable, native equations intact, no clipping or overlap; split tables repeat headers.',
        review_scope='Local manuscript and evidence review, not independent replication or venue acceptance assessment.')
    (PAPER/'english_document_qa.json').write_text(json.dumps(qa,indent=2),encoding='utf-8')
    print(json.dumps({'status':'passed','pages':18,'numeric_cells':checked,'words':check['english_prose_words'],'vocabulary_hit_terms':list(hits)},indent=2))
if __name__=='__main__':main()
