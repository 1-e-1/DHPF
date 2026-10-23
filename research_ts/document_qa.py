"""Record final structural checks and the completed human-visible page review."""
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path
from docx import Document
import pymupdf

ROOT = Path(__file__).resolve().parents[1]

def main():
    source = ROOT / 'DHPF_时序数据生成研究论文.docx'
    doc = Document(source)
    equations = doc._element.xpath('.//m:oMath')
    assert len(equations) == 11
    assert all(list(e.itertext()) for e in equations)
    assert len(doc.tables) == 7 and len(doc.inline_shapes) == 5
    assert '{{' not in '\n'.join(p.text for p in doc.paragraphs)
    frozen = json.loads((ROOT/'results/timeseries_v2/frozen_original.json').read_text(encoding='utf-8'))
    for name, expected in frozen['sha256'].items():
        assert hashlib.sha256((ROOT/name).read_bytes()).hexdigest() == expected, name
    render = ROOT/'tmp/paper_final_render'
    with pymupdf.open(next(render.glob('*.pdf'))) as pdf:
        assert len(pdf) == 15
        assert all(len(page.get_text()) > 100 for page in pdf)
    record = dict(
        checked_utc=datetime.now(timezone.utc).isoformat(),
        document=str(source.relative_to(ROOT)),
        sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
        pages=15, native_editable_equations=11, tables=7, figures=5,
        visual_review=dict(status='pass', pages_reviewed=list(range(1,16)),
            reviewer='primary research agent, visual inspection of final 120 dpi page renders',
            checks=['formula glyphs and completeness','text and table bounds',
                    'repeated table headers','figure and caption readability',
                    'page numbering and no trailing orphan page']),
        renderer='documents/render_docx.py with local official LibreOffice 25.2.7.2 and PyMuPDF raster adapter',
        original_source_files_unchanged=len(frozen['sha256']),
        scientific_claim_status='partial; local empirical advantages, not overall superiority',
        scientific_audit='results/timeseries_v2/EXPERIMENT_AUDIT.json',
    )
    (ROOT/'paper/document_qa.json').write_text(json.dumps(record,ensure_ascii=False,indent=2),encoding='utf-8')
    tracker=ROOT/'refine-logs/EXPERIMENT_TRACKER.md'
    text=tracker.read_text(encoding='utf-8').replace('完整正文、11公式、7表、5图；视觉验收见QA记录',
        '完成：15页、11公式、7表、5图；15页逐页视觉验收通过，见 paper/document_qa.json')
    tracker.write_text(text,encoding='utf-8')
    print(json.dumps(record,ensure_ascii=False,indent=2))

if __name__=='__main__':
    main()
