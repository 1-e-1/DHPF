"""Build the complete Chinese manuscript from verified result artifacts.

Document design: narrative_proposal, memo_masthead title-stack pattern;
named ChineseAcademic override: A4, 22 mm margins, SimSun/Times New Roman,
black headings, compact scholarly title, native editable OMML equations.
"""
import copy,csv,json,re,hashlib
from pathlib import Path
import numpy as np
from docx import Document
from docx.shared import Mm,Pt,RGBColor,Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from paper_fonts import normalize
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'results/timeseries_v2';PAPER=ROOT/'paper'
FINAL=ROOT/'DHPF_时序数据生成研究论文.docx'
MODELS=['gan','ddpm_v','icfm','otcfm','dhpf','drifting']
LABELS={'gan':'GAN','ddpm_v':'DDPM-v','icfm':'I-CFM','otcfm':'OT-CFM','dhpf':'DHPF','drifting':'Drifting-raw'}
def read(p):return json.loads(p.read_text(encoding='utf-8'))
def fmt(r,k,digits=4,newline=False):return f"{r[k+'_mean']:.{digits}f}"+('\n±' if newline else '±')+f"{r[k+'_std']:.{digits}f}"
def stylefont(style,name='Times New Roman',east='宋体',size=11):
    style.font.name=name;style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
    fonts=style.element.get_or_add_rPr().get_or_add_rFonts();fonts.set(qn('w:eastAsia'),east);fonts.set(qn('w:ascii'),name);fonts.set(qn('w:hAnsi'),name)
def setstyle(doc,name,size,before=0,after=6,line=1.25,bold=False):
    st=doc.styles[name] if name in doc.styles else doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    stylefont(st,size=size);st.font.bold=bold
    pf=st.paragraph_format;pf.space_before=Pt(before);pf.space_after=Pt(after);pf.line_spacing=line;pf.widow_control=True
    return st
def table(doc,caption,headers,rows):
    p=doc.add_paragraph(caption,'Table Caption');p.paragraph_format.keep_with_next=True
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    total=9411;first=1750 if len(headers)>4 else 2200; widths=[first]+[(total-first)//(len(headers)-1)]*(len(headers)-1);widths[-1]+=total-sum(widths)
    pr=t._tbl.tblPr
    for tag in ['tblW','tblInd','tblLayout','tblCellMar','tblBorders']:
        old=pr.find(qn('w:'+tag))
        if old is not None:pr.remove(old)
    for tag,attrs in [('tblW',{'w':str(total),'type':'dxa'}),('tblInd',{'w':'120','type':'dxa'}),('tblLayout',{'type':'fixed'})]:
        el=OxmlElement('w:'+tag)
        for k,v in attrs.items():el.set(qn('w:'+k),v)
        pr.append(el)
    margins=OxmlElement('w:tblCellMar')
    for side,val in [('top',80),('bottom',80),('start',120),('end',120)]:
        e=OxmlElement('w:'+side);e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa');margins.append(e)
    pr.append(margins)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'C0C0C0');borders.append(e)
    pr.append(borders)
    for col,w in zip(t.columns,widths):col.width=Twips(w)
    for cells,values in [(t.rows[0].cells,headers)]+[(t.add_row().cells,vals) for vals in rows]:
        for cell,val,w in zip(cells,values,widths):
            cell.width=Twips(w);cell.vertical_alignment=1;cell.text=str(val)
            for p in cell.paragraphs:p.style=doc.styles['Table Text'];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=False
    for cell in t.rows[0].cells:
        for p in cell.paragraphs:
            for run in p.runs:run.bold=True
        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F4F6F9');cell._tc.get_or_add_tcPr().append(shade)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in t.rows:
        cannot=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cannot)
    if len(rows)<=3:
        for row in t.rows[:-1]:
            for cell in row.cells:
                for p in cell.paragraphs:p.paragraph_format.keep_with_next=True
    else:
        for cell in t.rows[-1].cells:
            for p in cell.paragraphs:p.paragraph_format.keep_with_next=True
        doc.add_paragraph('数值为均值±样本标准差；主方法 n=5 个训练种子。', 'Table Note')
    return t

def main():
    assert read(OUT/'analysis_status.json')['complete'], 'Main evidence incomplete'
    s={(r['dataset'],r['model'],r['variant']):r for r in read(OUT/'summary.json')}
    refs=read(OUT/'real_reference.json')
    ridge={}
    for m in MODELS:
        name='ddpm' if m=='ddpm_v' else m
        vals=[read(OUT/'ridge_probe'/f'har__{name}__base__seed{k}.json')['test_macro_f1'] for k in range(5)]
        ridge[m]={'f1_mean':float(np.mean(vals)),'f1_std':float(np.std(vals,ddof=1))}
    rr=read(OUT/'ridge_probe/real_matched.json')['test_macro_f1']
    dh=s['har','dhpf','base'];gan=s['har','gan','base'];ic=s['har','icfm','base']
    text=(PAPER/'manuscript_template.txt').read_text(encoding='utf-8')
    replacements={
        'har_f1':fmt(dh,'synthetic_test_macro_f1'),
        'har_real_f1':f"{dh['real_matched_test_macro_f1_mean']:.4f}",
        'har_results':f"HAR 条件平均结果见表4。DHPF 的 SW₂² 为{fmt(dh,'sw2_squared')}，ACF 误差为{fmt(dh,'autocorrelation_rmse')}，Corr₀ 误差为{fmt(dh,'cross_correlation_rmse')}。在这些分布与依赖指标上，其表现没有延续合成任务中的局部领先：GAN 的 SW₂² 为{gan['sw2_squared_mean']:.4f}，且 GAN、DDPM-v 和两种 CFM 的 ACF、Corr₀ 均值均低于 DHPF。DHPF 的逻辑回归 TSTR macro-F1 为{fmt(dh,'synthetic_test_macro_f1')}，GAN 为{fmt(gan,'synthetic_test_macro_f1')}，I-CFM 为{fmt(ic,'synthetic_test_macro_f1')}。改用 ridge 分类器后，DHPF 为{fmt(ridge['dhpf'],'f1')}，I-CFM 为{fmt(ridge['icfm'],'f1')}，等量真实数据为{rr:.4f}。两种探针和条件分布指标共同表明，真实时序上的效用需要单独评价，当前结果不支持 DHPF 在 HAR 上全面领先。",
        'distribution_diagnostics':f"进一步的幅值诊断显示，DHPF 在 Sine/Lorenz 上的通道方差与真实测试数据之比分别为{s['sine','dhpf','base']['variance_ratio_mean']:.3f}和{s['lorenz','dhpf','base']['variance_ratio_mean']:.3f}，一阶差分能量之比分别为{s['sine','dhpf','base']['difference_energy_ratio_mean']:.2f}和{s['lorenz','dhpf','base']['difference_energy_ratio_mean']:.2f}。这些结果解释了为何归一化相关指标较低仍可能伴随幅值和波形质量的不足。训练数据子样本与测试集之间的 SW₂² 参照均值仅为{refs['sine']['sw2_squared']['mean']:.4f}和{refs['lorenz']['sw2_squared']['mean']:.4f}，说明 DHPF 与真实分布仍存在明显差距。这一参照包含抽样误差，HAR 中还包含受试者分布偏移，不作为数学下界。",
        'efficiency_results':f"从训练更新的同步计时看，Sine 上 DHPF、I-CFM 与 OT-CFM 分别约为{1000*s['sine','dhpf','base']['update_seconds_mean']:.1f}、{1000*s['sine','icfm','base']['update_seconds_mean']:.1f}和{1000*s['sine','otcfm','base']['update_seconds_mean']:.1f} ms/update。DHPF 需要积分反传和终点指派，当前实现没有训练速度优势。GAN 和 Drifting-raw 推理只调用一次网络，CFM 与 DHPF 使用64次场查询，DDPM-v 使用1000次去噪；由于单次查询的成本不同，NFE 不能直接换算为相同硬件上的速度倍数。推理计时使用同一设备、固定批量与同步测量独立归档，作为实现层面的诊断。"
    }
    replacements['har_results']+=f"Drifting-raw 的逻辑回归 F1 为{fmt(s['har','drifting','base'],'synthetic_test_macro_f1')}，但其 SW₂² 为{s['har','drifting','base']['sw2_squared_mean']:.4f}；较高分类效用也没有对应较好的完整分布拟合。"
    for key,val in replacements.items():text=text.replace('{{'+key+'}}',val)
    (PAPER/'manuscript_resolved.txt').write_text(text,encoding='utf-8')
    doc=Document();sec=doc.sections[0];sec.page_width=Mm(210);sec.page_height=Mm(297)
    for side in ['top_margin','bottom_margin','left_margin','right_margin']:setattr(sec,side,Mm(22))
    sec.header_distance=Mm(10);sec.footer_distance=Mm(12)
    normal=setstyle(doc,'Normal',11,after=5);normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;normal.paragraph_format.first_line_indent=Pt(22)
    setstyle(doc,'Title',18,after=6,bold=True).paragraph_format.first_line_indent=Pt(0)
    for border in doc.styles['Title'].element.xpath('./w:pPr/w:pBdr'):
        border.getparent().remove(border)
    setstyle(doc,'Subtitle',11,after=12).paragraph_format.first_line_indent=Pt(0)
    for name,size,bef,aft in [('Heading 1',14,14,8),('Heading 2',12,10,6),('Heading 3',11,8,4)]:
        st=setstyle(doc,name,size,bef,aft,bold=True);st.paragraph_format.first_line_indent=Pt(0);st.paragraph_format.keep_with_next=True
    for name,size,after in [('Caption',9,8),('Table Caption',10,5),('Table Note',9,4),('Table Text',9,0),('Equation',10,8),('Reference',9.5,5)]:
        st=setstyle(doc,name,size,after=after,line=1.12 if 'Table' in name else 1.2);st.paragraph_format.first_line_indent=Pt(0)
    doc.styles['Table Note'].paragraph_format.space_before=Pt(4)
    doc.styles['Reference'].paragraph_format.left_indent=Pt(18);doc.styles['Reference'].paragraph_format.first_line_indent=Pt(-18)
    header=sec.header.paragraphs[0];header.text='DHPF · 多变量时序生成';header.paragraph_format.first_line_indent=Pt(0)
    for r in header.runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('666666')
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT;footer.paragraph_format.first_line_indent=Pt(0)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    # Reuse the project's already corrected, native Word mathematics.
    source=Document(ROOT/'DHPF_V3_基础论文_数学修订版.docx');equations={}
    for t in source.tables:
        tag=t.cell(0,len(t.columns)-1).text
        if re.fullmatch(r'\(\d+\)',tag):equations[int(tag[1:-1])]=t._tbl.xpath('.//m:oMath')
    eqcount=0
    def add_equation(old):
        nonlocal eqcount
        eqcount+=1;p=doc.add_paragraph(style='Equation');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        assert equations[old],f'Equation {old} is empty'
        for math in equations[old]:
            math=copy.deepcopy(math)
            # Portable notation avoids unsupported supplementary-plane glyphs
            # and LibreOffice interpreting the parallel sign as slanted bars.
            for node in math.iter(qn('m:t')):
                if node.text in {'𝒩','𝔖','∥'}:
                    node.text={'𝒩':'N','𝔖':'S','∥':'||'}[node.text]
                    run=node.getparent();pr=run.find(qn('m:rPr'))
                    if pr is not None:run.remove(pr)
                    pr=OxmlElement('m:rPr');pr.append(OxmlElement('m:nor'));run.insert(0,pr)
            p._p.append(math)
        p.add_run(f'    ({eqcount})');p.paragraph_format.keep_together=True
    def main_table(dataset):
        keys=['sw2_squared','power_rmse','autocorrelation_rmse','cross_correlation_rmse','synthetic_test_forecast_mse']
        rows=[[LABELS[m]]+[fmt(s[dataset,m,'base'],k,newline=True) for k in keys] for m in MODELS]
        real=[f"{refs[dataset][k]['mean']:.4f}\n±{refs[dataset][k]['std']:.4f}" for k in keys[:-1]]+[f"{s[dataset,'dhpf','base']['real_matched_test_forecast_mse_mean']:.4f}"]
        rows.append(['Real→Real']+real)
        table(doc,'表'+('2' if dataset=='sine' else '3')+'　'+dataset.title()+' 的完整窗口生成结果（均越低越好）',['方法','SW₂²','谱 RMSE','ACF RMSE','Corr₀ RMSE','TSTR MSE'],rows)
        doc.add_paragraph('Real→Real 的前四项为五次真实训练子抽样对测试集的参照，误差棒来源不是训练种子；TSTR 列为固定等量真实数据训练的评价器结果。','Table Note')
    def add_table(key):
        if key=='datasets':table(doc,'表1　时序生成任务与数据规模',['任务','窗口形状','训练 / 验证 / 测试','生成条件'],[['Sine','128×3','10000 / 2000 / 2000','无条件'],['Lorenz','128×3','10000 / 2000 / 2000','无条件'],['UCI HAR','128×9','5883 / 1469 / 2947','6类活动']])
        elif key in ['sine','lorenz']:main_table(key)
        elif key=='har':
            rows=[[LABELS[m]]+[fmt(s['har',m,'base'],k,newline=True) for k in ['sw2_squared','autocorrelation_rmse','cross_correlation_rmse','synthetic_test_macro_f1']]+[fmt(ridge[m],'f1',newline=True)] for m in MODELS]
            rows.append(['Real→Real']+[f"{refs['har'][k]['mean']:.4f}\n±{refs['har'][k]['std']:.4f}" for k in ['sw2_squared','autocorrelation_rmse','cross_correlation_rmse']]+[f"{dh['real_matched_test_macro_f1_mean']:.4f}",f'{rr:.4f}'])
            table(doc,'表4　HAR 类别条件生成与 TSTR 分类',['方法','SW₂² ↓','ACF ↓','Corr₀ ↓','LR F1 ↑','Ridge F1 ↑'],rows)
            doc.add_paragraph('分布与依赖指标为六类等权平均；F1 是全部测试受试者的 macro-F1。Real→Real 前三项是训练子抽样参照，不是理论下界。','Table Note')
        elif key=='extra_metrics':
            rows=[[LABELS[m]]+[fmt(s[d,m,'base'],'mmd2_unbiased',newline=True) for d in ['sine','lorenz','har']]+[fmt(s['har',m,'base'],'power_rmse',newline=True)] for m in MODELS]
            table(doc,'表A1　补充的分布与频谱指标（越低越好）',['方法','Sine MMD²','Lorenz MMD²','HAR MMD²','HAR 谱 RMSE'],rows)
        elif key=='har_classes':
            with (OUT/'per_class.csv').open(encoding='utf-8-sig') as f:raw=list(csv.DictReader(f))
            rows=[]
            for c in range(6):
                records=[r for r in raw if r['dataset']=='har' and r['model']=='dhpf' and r['variant']=='base' and int(r['condition'])==c]
                assert len(records)==5
                rows.append([['行走','上楼','下楼','坐','站','躺'][c]]+[f"{np.mean([float(r[k]) for r in records]):.4f}\n±{np.std([float(r[k]) for r in records],ddof=1):.4f}" for k in ['sw2_squared','autocorrelation_rmse','cross_correlation_rmse','power_rmse']])
            table(doc,'表A2　DHPF 在 HAR 各类活动上的结果',['活动','SW₂²','ACF RMSE','Corr₀ RMSE','谱 RMSE'],rows)
        elif key=='ablations':
            variants=[('dhpf','base','完整 DHPF'),('dhpf','fixed_full','固定中心'),('dhpf','radial','仅径向'),('dhpf','tangential','仅切向'),('direct','base','MLP 加速度'),('augmented','base','扩展一阶 ODE')]
            rows=[[label]+[fmt(s['sine',m,v],k,newline=True) for k in ['sw2_squared','autocorrelation_rmse','cross_correlation_rmse','synthetic_test_forecast_mse']] for m,v,label in variants]
            table(doc,'表B1　Sine 结构与机制对照',['模型','SW₂²','ACF RMSE','Corr₀ RMSE','TSTR MSE'],rows)
        else:raise ValueError(key)
    blocks=text.strip().split('\n\n')
    for i,block in enumerate(blocks):
        if block.startswith('{{equation:'):add_equation(int(block[len('{{equation:'):-2]))
        elif block.startswith('{{table:'):add_table(block[len('{{table:'):-2])
        elif block.startswith('{{figure:'):
            name=block[len('{{figure:'):-2];p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.keep_with_next=True;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(PAPER/'figures'/(name+'.png')),width=Mm(164))
            for d in p._p.xpath('.//wp:docPr'):d.set('descr',name.replace('_',' '))
        elif block.startswith('### '):doc.add_paragraph(block[4:],'Heading 2')
        elif block.startswith('## '):doc.add_paragraph(block[3:],'Heading 1')
        elif block.startswith('# '):doc.add_paragraph(block[2:],'Title').alignment=WD_ALIGN_PARAGRAPH.CENTER
        elif i==1:doc.add_paragraph(block,'Subtitle').alignment=WD_ALIGN_PARAGRAPH.CENTER
        elif re.match(r'^图\d',block):doc.add_paragraph(block,'Caption')
        elif re.match(r'^\[\d+\]',block):doc.add_paragraph(block,'Reference')
        else:doc.add_paragraph(block)
    doc.core_properties.title='DHPF：面向多变量时序生成的共享场源二阶动力学方法';doc.core_properties.subject='完整方法、五种子时序对照与证据边界';doc.core_properties.author='';doc.core_properties.comments=''
    normalize(doc)
    doc.save(FINAL)
    assert '{{' not in '\n'.join(p.text for p in doc.paragraphs)
    assert len(doc._element.xpath('.//m:oMath'))>=eqcount
    (PAPER/'design_tokens.json').write_text(json.dumps(dict(preset='narrative_proposal',header_pattern='memo_masthead',named_override='ChineseAcademic',page_mm=[210,297],margins_mm=22,table_width_dxa=9411,body_size_pt=11,body_after_pt=5,body_line_spacing=1.25,fonts=['Times New Roman','宋体','Cambria Math'],equations=eqcount,author_identity='omitted until supplied'),ensure_ascii=False,indent=2),encoding='utf-8')
    print(FINAL);print('paragraphs',len(doc.paragraphs),'tables',len(doc.tables),'equations',eqcount)

if __name__=='__main__':main()
